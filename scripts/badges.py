#!/usr/bin/env -S uv run
# /// script
# requires-python = ">=3.11"
# dependencies = [
#     "requests>=2.32",
#     "openpyxl>=3.1",
#     "PyYAML>=6.0",
# ]
# ///
"""
BARGE 2026 badge pipeline — round 1: canonical attendee list.

Fetches the current Zeffy guest list (or reads a saved XLSX), collapses
multi-purchase rows into one row per attendee, keeps the extra fields
badges need (last name, hometown / city+state), and emits:

  1. A canonical CSV to --output (default scripts/badges_canonical.csv)
     — one row per attendee, ordered by their earliest ticket number.
  2. A validation report to stdout — missing nicknames, duplicate
     nicknames across different people, likely overflow risks, and a
     list of attendees missing a hometown value.

Round 2 will consume the same canonical list to render Avery 74459
PDFs, but that's not built yet.

Reuses the Zeffy client from zeffy_poll.py — same cookies, same
Cloudflare-friendly HTTP headers, same XLSX endpoint.  Unlike the
poller this tool keeps "Name For Badge", "Nickname For Badge",
"City/State For Badge", parses a last name off the display name, and
performs a case-insensitive dedup so e.g. "MRSTCAO" and "Mrstcao"
merge into one row.

Usage:
    uv run scripts/badges.py --campaign <UUID>
    uv run scripts/badges.py --campaign <UUID> --last-export scripts/last_export.xlsx
    uv run scripts/badges.py --campaign <UUID> --output /tmp/canonical.csv
"""

import argparse
import csv
import io
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# Reuse the poller's Zeffy client.  Importing by module works because
# scripts/ is on sys.path when running via `uv run scripts/badges.py`.
from zeffy_poll import (  # noqa: E402
    AuthExpired,
    fetch_xlsx_bytes,
    format_ticket_type_flags,
    load_session,
    parse_ticket_number,
    ticket_type_to_flags,
)


# Rough thresholds for the "will this overflow the badge?" flags.
# Tune once we've actually printed a sheet — for now, err on the
# side of surfacing things worth eyeballing.
NICKNAME_OVERFLOW = 20
DISPLAY_NAME_OVERFLOW = 25


@dataclass
class Attendee:
    display_name: str
    last_name: str
    nickname: str
    hometown: str
    tickets: list[int] = field(default_factory=list)
    ticket_type: str = ""


def find_col(header: list[str], *candidates: str) -> int | None:
    for cand in candidates:
        for i, h in enumerate(header):
            if cand in h:
                return i
    return None


def parse_extended(xlsx_bytes: bytes) -> list[Attendee]:
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True,
                                data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(c).strip().lower() if c else "" for c in rows[0]]

    badge_col = find_col(header, "name for badge", "badge name")
    guest_col = find_col(header, "guest")
    buyer_col = find_col(header, "buyer")
    ticket_col = find_col(header, "ticket number")
    type_col = find_col(header, "ticket type")
    nickname_col = find_col(header, "nickname for badge", "nickname",
                            "questions", "notes")
    city_col = find_col(header, "city/state for badge", "city/state",
                        "hometown", "city")

    if ticket_col is None or (badge_col is None and guest_col is None
                              and buyer_col is None):
        sys.exit(f"Could not find required columns in header: {header}")

    def cell(row: tuple, idx: int | None) -> str:
        if idx is None or idx >= len(row) or row[idx] is None:
            return ""
        return str(row[idx]).strip()

    raw = []
    for r in rows[1:]:
        if not r or all(c is None for c in r):
            continue
        display_name = (cell(r, badge_col) or cell(r, guest_col)
                        or cell(r, buyer_col))
        if not display_name:
            continue
        parts = display_name.split()
        last_name = parts[-1] if len(parts) > 1 else ""
        raw.append({
            "display_name": display_name,
            "last_name": last_name,
            "nickname": cell(r, nickname_col),
            "hometown": cell(r, city_col),
            "ticket": parse_ticket_number(cell(r, ticket_col)),
            "type_flags": ticket_type_to_flags(cell(r, type_col)),
        })

    # Sort by ticket ascending so the earliest row wins for display fields.
    raw.sort(key=lambda r: (r["ticket"] is None, r["ticket"] or 0))

    # Case-insensitive dedup on (display_name, nickname).  Values kept
    # verbatim from the earliest row; later rows contribute additional
    # tickets, a hometown backfill if the earlier row had none, and
    # union'd ticket-type flags.
    by_key: dict[tuple[str, str], Attendee] = {}
    flags_by_key: dict[tuple[str, str], set[str]] = {}
    for r in raw:
        key = (r["display_name"].casefold(), r["nickname"].casefold())
        if key not in by_key:
            by_key[key] = Attendee(
                display_name=r["display_name"],
                last_name=r["last_name"],
                nickname=r["nickname"],
                hometown=r["hometown"],
                tickets=[r["ticket"]] if r["ticket"] is not None else [],
            )
            flags_by_key[key] = set()
        else:
            a = by_key[key]
            if r["ticket"] is not None:
                a.tickets.append(r["ticket"])
            if not a.hometown and r["hometown"]:
                a.hometown = r["hometown"]
        flags_by_key[key] |= r["type_flags"]

    out = []
    for key, a in by_key.items():
        a.ticket_type = format_ticket_type_flags(flags_by_key[key])
        out.append(a)
    # Present sorted by earliest ticket for downstream ergonomics.
    out.sort(key=lambda a: min(a.tickets) if a.tickets else 9999)
    return out


def validation_report(attendees: list[Attendee]) -> None:
    problems: list[str] = []

    # Missing nickname — badge would have a blank headline.
    for a in attendees:
        if not a.nickname:
            tickets = ",".join(str(t) for t in a.tickets) if a.tickets else "?"
            problems.append(
                f"  missing nickname: '{a.display_name}' (ticket #{tickets})"
            )

    # Same nickname on two different display_names.
    by_nick: dict[str, list[Attendee]] = defaultdict(list)
    for a in attendees:
        if a.nickname:
            by_nick[a.nickname.casefold()].append(a)
    for group in by_nick.values():
        distinct_names = {a.display_name.casefold() for a in group}
        if len(distinct_names) > 1:
            names = ", ".join(f"'{a.display_name}'" for a in group)
            problems.append(
                f"  duplicate nickname '{group[0].nickname}' on: {names}"
            )

    # Overflow risk.  Thresholds are placeholders until we've printed.
    for a in attendees:
        if len(a.nickname) > NICKNAME_OVERFLOW:
            problems.append(
                f"  long nickname ({len(a.nickname)} chars): "
                f"'{a.nickname}' — {a.display_name}"
            )
        if len(a.display_name) > DISPLAY_NAME_OVERFLOW:
            problems.append(
                f"  long display name ({len(a.display_name)} chars): "
                f"'{a.display_name}'"
            )

    print(f"Validation: {len(problems)} flag(s).")
    for p in problems:
        print(p)

    missing_hometown = [a for a in attendees if not a.hometown]
    if missing_hometown:
        names = ", ".join(a.display_name for a in missing_hometown)
        print()
        print(f"Missing hometown ({len(missing_hometown)}): {names}")


def emit_csv(attendees: list[Attendee], path: Path) -> None:
    with path.open("w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["ticket", "display_name", "last_name", "nickname",
                    "hometown", "ticket_type", "all_tickets"])
        for a in attendees:
            primary = min(a.tickets) if a.tickets else ""
            w.writerow([
                primary,
                a.display_name,
                a.last_name,
                a.nickname,
                a.hometown,
                a.ticket_type,
                ",".join(str(t) for t in sorted(a.tickets)),
            ])


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--campaign", required=True,
                    help="Zeffy campaign UUID.")
    ap.add_argument(
        "--output", type=Path,
        default=Path(__file__).parent / "badges_canonical.csv",
        help="Where to write the canonical CSV.  Default: scripts/badges_canonical.csv",
    )
    ap.add_argument(
        "--last-export", type=Path,
        help="Skip the Zeffy fetch and read this XLSX file instead.  "
             "Useful when Zeffy is flaky or for offline iteration.",
    )
    args = ap.parse_args()

    if args.last_export:
        if not args.last_export.exists():
            sys.exit(f"No file at {args.last_export}")
        xlsx_bytes = args.last_export.read_bytes()
        print(f"Read {len(xlsx_bytes)} bytes from {args.last_export}.")
    else:
        session = load_session()
        try:
            xlsx_bytes = fetch_xlsx_bytes(session, args.campaign)
        except AuthExpired as e:
            sys.exit(
                f"Zeffy auth failed ({e}).  Refresh cookies via "
                "scripts/zeffy_login_headless.py."
            )
        print(f"Fetched {len(xlsx_bytes)} bytes from Zeffy.")

    attendees = parse_extended(xlsx_bytes)
    print(f"Parsed {len(attendees)} canonical attendee(s).")
    print()
    validation_report(attendees)
    print()

    emit_csv(attendees, args.output)
    print(f"Wrote {args.output}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
